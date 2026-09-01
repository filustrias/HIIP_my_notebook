"""
Gera o capítulo `planejamento-e-roadmap.qmd` a partir de HIIP_plano_completo.xlsx.

A fonte de verdade é a aba "2. Master" (frente > resultado > tarefa). As abas
"0. Legenda", "1. Frentes" e "4. Sprints" fornecem as convenções, a validação
externa e o calendário.

Uso:
    python scripts/gerar_roadmap.py

Reexecutar depois de editar a planilha reescreve o capítulo por completo. O
texto interpretativo de cada frente vive no dicionário NOTAS abaixo, para
sobreviver à regeneração — edite-o aqui, nunca no .qmd.
"""

import datetime
import re
import shutil
import subprocess
import tempfile
from collections import OrderedDict
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
XLSX = RAIZ / "HIIP_plano_completo.xlsx"
SAIDA = RAIZ / "planejamento-e-roadmap.qmd"


def abrir_planilha(caminho):
    """Abre o .xlsx mesmo com o Excel segurando um lock exclusivo.

    O Excel abre o ficheiro negando partilha de leitura, o que faz o open() do
    Python falhar com PermissionError. O CopyFile do Windows consegue ler à
    mesma, por isso caímos para uma cópia temporária.
    """
    try:
        return openpyxl.load_workbook(caminho, data_only=True)
    except PermissionError:
        pass

    tmp = Path(tempfile.gettempdir()) / f"_hiip_{caminho.name}"
    try:
        shutil.copy2(caminho, tmp)
    except PermissionError:
        subprocess.run(
            ["powershell", "-NoProfile", "-Command",
             f"Copy-Item -LiteralPath '{caminho}' -Destination '{tmp}' -Force"],
            check=True, capture_output=True,
        )
    print(f"  (planilha aberta noutro programa; li uma cópia de {tmp})")
    return openpyxl.load_workbook(tmp, data_only=True)

# --------------------------------------------------------------------------
# Texto autoral por frente. Preservado entre regenerações.
# --------------------------------------------------------------------------
NOTAS = {
    "TWG-PHC": (
        "O grupo técnico de trabalho para os cuidados de saúde primários é a "
        "instância que valida quase todo o resto do plano: sete das doze frentes "
        "têm a sua validação externa marcada numa sessão do TWG-PHC. Reconstituí-lo "
        "não é uma frente entre outras — é o pré-requisito que governa o calendário "
        "inteiro."
    ),
    "Background document - PHC": (
        "O diagnóstico. Herda trabalho de consultoras anteriores, que é auditado "
        "antes de ser reaproveitado, e combina uma espinha dorsal quantitativa com "
        "visitas piloto de campo. É o insumo direto do *strategy document*: nenhuma "
        "priorização estratégica pode ser defendida sem ele."
    ),
    "Strategy document - PHC": (
        "Converte o diagnóstico em escolhas, por uma cadeia explícita: mapa de "
        "barreiras e lacunas → teoria da mudança → áreas priorizadas. A validação "
        "exige que o MINSA confirme que a estratégia serve de documento de política "
        "habilitante — sem isso, o *investment plan* e o pacote essencial ficam sem "
        "base normativa."
    ),
    "HHFA e diagnóstico situacional": (
        "A frente mais longa do plano, com validação apenas em junho de 2027. "
        "Envolve negociação interagências com o UNICEF e o UNFPA, adesão a um "
        "inquérito já contratado, e um legado explícito: propor ao MINSA um "
        "mecanismo rotineiro de avaliação estrutural, para que a avaliação não "
        "termine com o projeto. O conteúdo do instrumento está tratado na parte "
        "HHFA deste caderno."
    ),
    "PFM": (
        "Gestão das finanças públicas. Corre com o seu próprio grupo de trabalho, "
        "o TWG-PFM, separado do TWG-PHC e com a sua própria validação. É a frente "
        "onde há mais trabalho já concluído."
    ),
    "Investment case - PHC": (
        "Único produto com **duas** validações externas: um diálogo político com a "
        "Ministra sobre a v1, em novembro de 2026, e uma validação técnica da v2 já "
        "com dados do HHFA, em fevereiro de 2027. A separação reconhece que "
        "convencer e comprovar são momentos distintos, com públicos e critérios "
        "distintos."
    ),
    "PHC investment plan (custeado)": (
        "Traduz a estratégia em números. A ordem importa: o consenso sobre as áreas "
        "de investimento precede qualquer custeio, porque custear antes de decidir "
        "onde investir produz um número que ninguém assume."
    ),
    "Essential package - PHC": (
        "Define o que os cuidados primários devem efetivamente oferecer, filtrando "
        "o UHC Compendium para o nível dos CSP e ajustando ao perfil epidemiológico "
        "do país."
    ),
    "Modelo de cuidados e perfis profissionais - PHC": (
        "A contraparte operacional do pacote essencial: se o pacote diz *o quê*, "
        "esta frente diz *por quem e como* — composição das equipas, funções "
        "efetivamente exercidas, perfis profissionais."
    ),
    "Project scoping - PHC": (
        "Mapeia quem já financia cuidados primários, onde e para quê, identifica "
        "sobreposições e lacunas, e converte o resultado em *project profiles* "
        "apresentáveis a financiadores."
    ),
    "Mobilização de financiamento - PHC": (
        "O fecho do arco: engajamento do BEI, de outros doadores e da delegação da "
        "UE, e um empréstimo do Banco Mundial. É a frente cuja validação está mais "
        "longe no tempo e que depende de praticamente todas as outras."
    ),
    "Relatórios contratuais do HIIP": (
        "Obrigações contratuais de reporte. É a única frente sem validação externa "
        "registada, porque o critério não é a aceitação de um cliente técnico mas "
        "o cumprimento do prazo."
    ),
}

CLASSE_ESTADO = {
    "Done": "st-done",
    "Feito": "st-done",
    "Feita": "st-done",
    "In progress": "st-prog",
    "Em curso": "st-prog",
    "Blocked": "st-block",
    "To do": "st-todo",
    "Backlog": "st-backlog",
    "Não iniciado": "st-none",
    "Não iniciada": "st-none",
}

MESES = [
    "janeiro", "fevereiro", "março", "abril", "maio", "junho",
    "julho", "agosto", "setembro", "outubro", "novembro", "dezembro",
]


def txt(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def data_legivel(iso):
    """2026-10-23 -> 23 de outubro de 2026"""
    if not iso:
        return ""
    try:
        d = datetime.date.fromisoformat(iso)
    except ValueError:
        return iso
    return f"{d.day} de {MESES[d.month - 1]} de {d.year}"


def celula(s):
    """Escapa o que quebraria uma tabela pipe do Markdown."""
    return txt(s).replace("|", "\\|").replace("\n", " ")


def badge(estado):
    if not estado:
        return ""
    return f"[{estado}]{{.{CLASSE_ESTADO.get(estado, 'st-none')}}}"


def titulo_callout(s):
    """Título de callout: sem aspas duplas, que fechariam o atributo."""
    return txt(s).replace('"', "'")


def chave_sprint(s):
    m = re.fullmatch(r"S(\d+)", s or "")
    return int(m.group(1)) if m else None


def intervalo_sprints(tarefas):
    nums = [n for n in (chave_sprint(t["Sprint"]) for t in tarefas) if n is not None]
    if not nums:
        return ""
    lo, hi = min(nums), max(nums)
    return f"S{lo}" if lo == hi else f"S{lo}–S{hi}"


def main():
    wb = abrir_planilha(XLSX)

    # ---- Master: fonte de verdade -------------------------------------
    ws = wb["2. Master"]
    hdr = [txt(c) for c in next(ws.iter_rows(min_row=3, max_row=3, values_only=True))]
    linhas = []
    for r in ws.iter_rows(min_row=4, values_only=True):
        d = {hdr[i]: txt(v) for i, v in enumerate(r) if i < len(hdr) and hdr[i]}
        if d.get("Tarefa") or d.get("Resultado (output)"):
            linhas.append(d)

    # ---- Legenda ------------------------------------------------------
    legenda = {}
    secao = None
    for row in wb["0. Legenda"].iter_rows(values_only=True):
        a, b = txt(row[0]), txt(row[1]) if len(row) > 1 else ""
        if a and not b:
            secao = a
            legenda[secao] = []
        elif a and b and secao:
            legenda[secao].append((a, b))

    # ---- Sprints ------------------------------------------------------
    sprints = []
    for row in wb["4. Sprints"].iter_rows(min_row=4, values_only=True):
        s = [txt(c) for c in row]
        if s and s[0]:
            sprints.append(s)

    # ---- Agrupar: frente > resultado > tarefas -------------------------
    frentes = OrderedDict()
    for ln in linhas:
        f = ln["Frente de trabalho (produto)"]
        frentes.setdefault(f, {"meta": ln, "outputs": OrderedDict()})
        frentes[f]["outputs"].setdefault(ln["Resultado (output)"], []).append(ln)

    total_tarefas = len(linhas)
    total_feitas = sum(1 for l in linhas if l["Estado tarefa"] == "Done")
    total_outputs = sum(
        len([o for o in d["outputs"] if o]) for d in frentes.values()
    )

    out = []
    A = out.append

    # =================== cabeçalho e enquadramento =====================
    A("# Planejamento e roadmap\n")
    A(
        "Este capítulo é a versão legível do plano de trabalho do HIIP. A fonte "
        "de verdade continua a ser a planilha `HIIP_plano_completo.xlsx`, na aba "
        "*Master*; o que está aqui é gerado a partir dela e serve para consultar "
        "e discutir, não para editar.\n"
    )
    A(
        f"O plano cobre **{len(frentes)} frentes de trabalho**, "
        f"**{total_outputs} resultados** e **{total_tarefas} tarefas**, "
        "distribuídas por 23 sprints quinzenais entre 3 de agosto de 2026 e "
        f"2 de julho de 2027. Neste momento há **{total_feitas} tarefas concluídas** "
        f"({total_feitas / total_tarefas:.1%}). Todas as tarefas estão atribuídas a AGJ.\n"
    )

    # ---- como ler -----------------------------------------------------
    A("## Como ler este plano\n")
    A(
        "O plano tem três níveis, e a distinção entre eles é o que o torna "
        "utilizável. Uma **frente de trabalho** produz um produto e é o único "
        "nível onde existe validação externa. Um **resultado** é uma entrega "
        "intermédia que fecha dentro de uma ou duas sprints, e é o único nível "
        "onde existe definição de feito. Uma **tarefa** é o trabalho concreto: "
        "não tem critério próprio, está feita ou não está.\n"
    )
    A(
        "A separação que mais importa é entre **definição de feito** e "
        "**validação externa**. A definição de feito reúne critérios internos, "
        "sob controlo de quem executa: o resultado fecha quando são cumpridos. "
        "A validação externa é a aceitação pelo cliente — o TWG-PHC, o TWG-PFM, "
        "a Ministra, os financiadores — tem data própria e **não está sob "
        "controlo de quem executa**. Um produto pode estar feito segundo todos "
        "os critérios internos e ainda assim não estar validado.\n"
    )
    A(
        "Algumas tarefas aparecem sem resultado associado. Não é omissão: são "
        "tarefas que servem diretamente a validação externa da frente, sem passar "
        "por uma entrega intermédia. Estão agrupadas à parte em cada frente.\n"
    )

    if "Definição de feito geral - aplica-se a todos os resultados" in legenda:
        A(
            '::: {.callout-important collapse="true" '
            'title="Definição de feito geral — aplica-se a todos os resultados"}\n'
        )
        A(
            "Estes critérios acumulam-se aos critérios próprios de cada resultado. "
            "Um resultado só fecha quando cumpre ambos.\n"
        )
        for _, d in legenda["Definição de feito geral - aplica-se a todos os resultados"]:
            A(f"1. {d}")
        A("\n:::\n")

    if "Estados" in legenda:
        A('::: {.callout-note collapse="true" title="Estados e convenções"}\n')
        A("| Nível | Estados possíveis |")
        A("|---|---|")
        for k, v in legenda["Estados"]:
            A(f"| {celula(k)} | {celula(v)} |")
        A("")
        if "Outras colunas" in legenda:
            A("| Coluna | Significado |")
            A("|---|---|")
            for k, v in legenda["Outras colunas"]:
                A(f"| {celula(k)} | {celula(v)} |")
            A("")
        A(":::\n")

    regra = legenda.get("Regra de promoção") or []
    texto_regra = regra[0][1] if regra else ""
    if not texto_regra:
        for row in wb["0. Legenda"].iter_rows(values_only=True):
            if len(row) > 1 and txt(row[1]).startswith("Aprendizado que muda"):
                texto_regra = txt(row[1])
    if texto_regra:
        A('::: {.callout-tip title="Regra de promoção"}')
        A(texto_regra)
        A(":::\n")

    # ---- panorama -----------------------------------------------------
    A("## Panorama das frentes\n")
    A(
        "Cada frente abaixo tem a sua secção, com a validação externa, os "
        "resultados e as tarefas. As datas de validação são o que estrutura o "
        "calendário: elas não se movem por conveniência interna.\n"
    )
    A("| Frente | Validação em | Resultados | Tarefas | Feitas | Sprints |")
    A("|:---|:---|---:|---:|---:|:---|")
    for nome, d in frentes.items():
        tarefas = [t for lst in d["outputs"].values() for t in lst]
        n_out = len([o for o in d["outputs"] if o])
        feitas = sum(1 for t in tarefas if t["Estado tarefa"] == "Done")
        dv = d["meta"]["Data validação"]
        A(
            f"| [{celula(nome)}](#{slug(nome)}) "
            f"| {data_legivel(dv) if dv else '—'} "
            f"| {n_out} | {len(tarefas)} | {feitas} "
            f"| {intervalo_sprints(tarefas) or '—'} |"
        )
    A("")

    # ---- calendário ---------------------------------------------------
    A('::: {.callout-note collapse="true" title="Calendário de sprints"}\n')
    A("Sprints quinzenais. A carga não é uniforme: as sprints S3 a S6 concentram ")
    A("o arranque de quase todas as frentes.\n")
    A("| Sprint | Início | Fim | Tarefas | Dias |")
    A("|:---|:---|:---|---:|---:|")
    for s in sprints:
        s = s + [""] * (5 - len(s))
        A(
            f"| {celula(s[0])} | {data_legivel(s[1]) or '—'} "
            f"| {data_legivel(s[2]) or '—'} | {celula(s[3])} | {celula(s[4])} |"
        )
    A("\n:::\n")

    # =================== uma secção por frente =========================
    A("## As frentes\n")

    for nome, d in frentes.items():
        meta = d["meta"]
        tarefas_todas = [t for lst in d["outputs"].values() for t in lst]
        feitas = sum(1 for t in tarefas_todas if t["Estado tarefa"] == "Done")
        n_out = len([o for o in d["outputs"] if o])

        A(f"### {nome} {{#{slug(nome)}}}\n")

        if nome in NOTAS:
            A(NOTAS[nome] + "\n")

        val = meta["Validação externa"]
        dv = meta["Data validação"]
        if val:
            A('::: {.callout-warning title="Validação externa"}')
            A(f"**Quando:** {data_legivel(dv)}  ")
            A(f"**Estado:** {badge(meta['Estado validação'])}\n")
            A(val)
            A(":::\n")
        else:
            A(
                "*Sem validação externa registada: o critério desta frente é o "
                "cumprimento do prazo contratual.*\n"
            )

        span = intervalo_sprints(tarefas_todas)
        n_t = len(tarefas_todas)
        frase = (
            f"{'É' if n_out == 1 else 'São'} {n_out} "
            f"{'resultado' if n_out == 1 else 'resultados'} e {n_t} "
            f"{'tarefa' if n_t == 1 else 'tarefas'}"
        )
        if span:
            frase += f", entre {span}"
        if feitas == 0:
            frase += ". Nenhuma concluída até agora."
        elif feitas == 1:
            frase += ". 1 concluída."
        else:
            frase += f". {feitas} concluídas."
        A(frase + "\n")

        for output, tarefas in d["outputs"].items():
            if output:
                dfeito = tarefas[0]["Definição de feito (resultado)"]
                A(
                    f'::: {{.callout-note collapse="true" '
                    f'title="{titulo_callout(output)}"}}\n'
                )
                A(
                    f"**Previsto para** {data_legivel(tarefas[0]['Data resultado']) or 'data por definir'} "
                    f"· **Estado** {badge(tarefas[0]['Estado resultado'])}\n"
                )
                if dfeito:
                    A(f"**Definição de feito.** {dfeito}\n")
            else:
                A(
                    '::: {.callout-note collapse="true" '
                    'title="Tarefas ligadas diretamente à validação da frente"}\n'
                )
                A(
                    "Tarefas sem resultado intermédio: servem diretamente a "
                    "validação externa desta frente.\n"
                )

            A("| Tarefa | Sprint | Est. | Depende de | Estado |")
            A("|:---|:---|---:|:---|:---|")
            for t in tarefas:
                est = t["Estimativa (dias)"]
                A(
                    f"| {celula(t['Tarefa'])} "
                    f"| {celula(t['Sprint']) or '—'} "
                    f"| {est + ' d' if est else '—'} "
                    f"| {celula(t['Depende de']) or '—'} "
                    f"| {badge(t['Estado tarefa'])} |"
                )
            A("")

            aprendizados = [
                (t["Tarefa"], t["Aprendizado"]) for t in tarefas if t.get("Aprendizado")
            ]
            if aprendizados:
                A("**Aprendizados registados**\n")
                for tarefa, ap in aprendizados:
                    A(f"- *{tarefa}* — {ap}")
                A("")

            A(":::\n")

    SAIDA.write_text("\n".join(out), encoding="utf-8")
    print(f"Escrito: {SAIDA}")
    print(
        f"  {len(frentes)} frentes, {total_outputs} resultados, "
        f"{total_tarefas} tarefas ({total_feitas} feitas)"
    )


def slug(s):
    s = s.lower()
    for a, b in zip("áàâãäéèêëíìîïóòôõöúùûüç", "aaaaaeeeeiiiiooooouuuuc"):
        s = s.replace(a, b)
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    return "frente-" + s


if __name__ == "__main__":
    main()
